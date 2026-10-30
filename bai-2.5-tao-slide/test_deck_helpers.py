import unittest
from pptx import Presentation
from pptx.util import Inches
from generate_deck import (
    add_flowchart_3_nodes,
    add_terminal_mockup,
    add_top_decorations
)

class TestDeckHelpers(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def test_add_top_decorations_has_10_slides_and_no_7_phut(self):
        add_top_decorations(self.slide, 1, 10)
        texts = [shape.text_frame.text for shape in self.slide.shapes if shape.has_text_frame]
        combined = " ".join(texts)
        self.assertIn("01 / 10", combined)
        self.assertIn("CHUYÊN ĐỀ CÔNG NGHỆ NỘI BỘ", combined)
        self.assertNotIn("7 PHÚT", combined)

    def test_add_flowchart_creates_3_nodes(self):
        initial_count = len(self.slide.shapes)
        nodes = [("Node 1", "Desc 1"), ("Node 2", "Desc 2"), ("Node 3", "Desc 3")]
        add_flowchart_3_nodes(self.slide, Inches(1), Inches(1), Inches(11.7), Inches(2.2), nodes)
        self.assertGreater(len(self.slide.shapes), initial_count + 2)

    def test_add_terminal_mockup(self):
        initial_count = len(self.slide.shapes)
        add_terminal_mockup(self.slide, Inches(1), Inches(1), Inches(5.5), Inches(4.5), "bash", "$ echo 1", ["[✓] Done"])
        self.assertGreater(len(self.slide.shapes), initial_count)

if __name__ == "__main__":
    unittest.main()
