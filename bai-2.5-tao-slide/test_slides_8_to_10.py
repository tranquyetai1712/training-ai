import unittest
from pptx import Presentation
from pptx.util import Inches
from generate_deck import build_slides_8_to_10

class TestSlides8To10(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        build_slides_8_to_10(self.prs)

    def test_slide_count_is_3(self):
        self.assertEqual(len(self.prs.slides), 3)

    def test_slide_8_has_value_pillars(self):
        s8 = self.prs.slides[0]
        text = " ".join([sh.text_frame.text for sh in s8.shapes if sh.has_text_frame])
        self.assertIn("nda", text.lower())
        self.assertIn("token", text.lower())

    def test_slide_9_has_terminal_command(self):
        s9 = self.prs.slides[1]
        text = " ".join([sh.text_frame.text for sh in s9.shapes if sh.has_text_frame])
        self.assertIn("ollama launch claude", text)

    def test_slide_10_has_checklist_and_qa(self):
        s10 = self.prs.slides[2]
        text = " ".join([sh.text_frame.text for sh in s10.shapes if sh.has_text_frame])
        self.assertIn("BẬT REMOTE CONTROL", text)
        self.assertIn("THẢO LUẬN", text)

if __name__ == "__main__":
    unittest.main()
