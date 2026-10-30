import unittest
from pptx import Presentation
from pptx.util import Inches
from generate_deck import build_slides_5_to_7

class TestSlides5To7(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        build_slides_5_to_7(self.prs)

    def test_slide_count_is_3(self):
        self.assertEqual(len(self.prs.slides), 3)

    def test_slide_5_has_flowchart_and_zero_trust(self):
        s5 = self.prs.slides[0]
        text = " ".join([sh.text_frame.text for sh in s5.shapes if sh.has_text_frame])
        self.assertIn("Mã QR", text)
        self.assertIn("Zero Trust", text)

    def test_slide_6_has_prevent_sleep_and_toggle(self):
        s6 = self.prs.slides[1]
        text = " ".join([sh.text_frame.text for sh in s6.shapes if sh.has_text_frame])
        self.assertIn("Prevent Sleep", text)
        self.assertIn("Kernel", text)

    def test_slide_7_has_menu_bar_and_notifications(self):
        s7 = self.prs.slides[2]
        text = " ".join([sh.text_frame.text for sh in s7.shapes if sh.has_text_frame])
        self.assertIn("Menu Bar", text)
        self.assertIn("Notifications", text)

if __name__ == "__main__":
    unittest.main()
