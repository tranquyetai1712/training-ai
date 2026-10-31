import unittest
from pptx import Presentation
from pptx.util import Inches
from generate_deck import build_slides_1_to_4

class TestSlides1To4(unittest.TestCase):
    def setUp(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        build_slides_1_to_4(self.prs)

    def test_slide_count_is_4(self):
        self.assertEqual(len(self.prs.slides), 4)

    def test_slide_1_has_hero_metrics(self):
        s1 = self.prs.slides[0]
        text = " ".join([sh.text_frame.text for sh in s1.shapes if sh.has_text_frame])
        self.assertIn("Tối Ưu Hóa Trải Nghiệm Phát Triển Với AI", text)
        self.assertIn("Multi-Device Control", text)
        self.assertNotIn("7 PHÚT", text.upper())

    def test_slide_2_has_4_pain_points_with_complete_sentences(self):
        s2 = self.prs.slides[1]
        text = " ".join([sh.text_frame.text for sh in s2.shapes if sh.has_text_frame])
        self.assertIn("Máy ngủ gật (Sleep Mode):", text)
        self.assertIn("Bị xích vào bàn làm việc:", text)
        self.assertIn("Đóng nhầm cửa sổ:", text)
        self.assertIn("Nguy cơ lộ mã nguồn dự án:", text)

    def test_slide_3_has_comparison(self):
        s3 = self.prs.slides[2]
        text = " ".join([sh.text_frame.text for sh in s3.shapes if sh.has_text_frame])
        self.assertIn("Gò bó vị trí địa lý:", text)
        self.assertIn("Tự do di chuyển tuyệt đối:", text)

    def test_slide_4_has_remote_control_definition(self):
        s4 = self.prs.slides[3]
        text = " ".join([sh.text_frame.text for sh in s4.shapes if sh.has_text_frame])
        self.assertIn("Không phải Remote Desktop:", text)
        self.assertIn("Máy trạm giữ quyền điều phối:", text)

if __name__ == "__main__":
    unittest.main()
